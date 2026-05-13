#!/usr/bin/env python3
"""
详细提取表1到表8的完整结构
"""

import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Error: python-docx not installed")
    sys.exit(1)

def extract_full_table(table, table_name):
    """提取完整表格内容"""
    print(f"\n{'='*80}")
    print(f"{table_name}")
    print(f"{'='*80}")
    print(f"尺寸: {len(table.rows)} 行 x {len(table.columns)} 列")
    
    if not table.rows:
        return
    
    # 提取所有数据
    all_data = []
    for row in table.rows:
        row_data = [cell.text.strip() for cell in row.cells]
        all_data.append(row_data)
    
    # 打印表头
    header = all_data[0] if all_data else []
    print(f"\n表头: {' | '.join(header)}")
    
    # 打印所有行
    print(f"\n完整数据（共{len(all_data)-1}行）:")
    for i, row_data in enumerate(all_data[1:], 1):
        print(f"行{i}: {' | '.join(row_data)}")
    
    return all_data

if __name__ == "__main__":
    file_path = "/workspace/projects/assets/宁波银行客户洞察及相关商机分析报告（完备版） _Kimi生成_20260303_发布后优化版20260310.docx"
    
    doc = Document(file_path)
    
    # 提取前8个表格（对应第二章的8个表格）
    table_names = [
        "表1：宁波银行主体及资质信息表",
        "表2：宁波银行规模与行业地位信息表", 
        "表3：宁波银行资产质量指标表",
        "表4：宁波银行资本充足率指标表",
        "表5：宁波银行盈利能力与行业对比表",
        "表6：按业务类型划分的业务收入及利润结构",
        "表7：按利息和非利息划分的收入和利润结构",
        "表8：综合经营三年发展趋势"
    ]
    
    print("提取第二章的8个表格...")
    print(f"文档共{len(doc.tables)}个表格")
    
    for i in range(min(8, len(doc.tables))):
        extract_full_table(doc.tables[i], table_names[i])
