#!/usr/bin/env python3
"""
解析Word文档，提取报告结构和内容
用于理解银行客户洞察报告的格式要求
"""

import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Error: python-docx not installed")
    sys.exit(1)

def parse_docx(file_path: str):
    """解析Word文档并提取结构化内容"""
    doc = Document(file_path)
    
    print("=" * 80)
    print("Word文档结构分析")
    print("=" * 80)
    
    # 提取段落
    print("\n【段落内容】")
    for i, para in enumerate(doc.paragraphs[:50]):  # 只显示前50个段落
        if para.text.strip():
            style = para.style.name if para.style else "Normal"
            print(f"{i+1}. [{style}] {para.text[:100]}")
    
    # 提取表格
    print(f"\n【表格信息】共发现 {len(doc.tables)} 个表格")
    for i, table in enumerate(doc.tables[:5]):  # 只显示前5个表格
        print(f"\n表格 {i+1}: {len(table.rows)} 行 x {len(table.columns)} 列")
        # 显示表头
        if table.rows:
            header = [cell.text[:20] for cell in table.rows[0].cells]
            print(f"  表头: {' | '.join(header)}")
    
    # 提取标题层级
    print("\n【标题结构】")
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            print(f"  {para.style.name}: {para.text[:80]}")
    
    return doc

if __name__ == "__main__":
    file_path = "./宁波银行客户洞察及相关商机分析报告（完备版）_Kimi生成_20260303_发布后优化版20260310.docx"
    
    if not Path(file_path).exists():
        print(f"Error: 文件不存在 - {file_path}")
        sys.exit(1)
    
    try:
        doc = parse_docx(file_path)
        print("\n解析完成！")
    except Exception as e:
        print(f"Error: 解析失败 - {str(e)}")
        sys.exit(1)
