#!/usr/bin/env python3
"""
银行客户洞察报告生成器
从JSON格式的内容生成符合规范的Word文档
"""

import sys
import json
import argparse
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Any

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("错误: 缺少 python-docx 库，请执行: pip install python-docx==1.1.2")
    sys.exit(1)


class BankReportGenerator:
    """银行客户洞察报告生成器"""

    def __init__(self, bank_name: str, period: str):
        """
        初始化生成器

        Args:
            bank_name: 银行名称
            period: 报告周期（如"2025年度"）
        """
        self.bank_name = bank_name
        self.period = period
        self.doc = Document()

    def _setup_styles(self):
        """设置文档样式"""
        # 设置默认字体
        self.doc.styles['Normal'].font.name = '宋体'
        self.doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        self.doc.styles['Normal'].font.size = Pt(10.5)

        # 设置标题样式
        for level in range(1, 5):
            style_name = f'Heading {level}'
            if style_name in self.doc.styles:
                heading_style = self.doc.styles[style_name]
                heading_style.font.name = '黑体'
                heading_style.font.bold = True
                heading_style.font.size = Pt(16 - level * 1.5)
                heading_style.paragraph_format.space_before = Pt(12)
                heading_style.paragraph_format.space_after = Pt(6)

    def _add_cover_page(self):
        """添加封面页"""
        # 银行名称
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(self.bank_name)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.name = '黑体'

        # 报告标题
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run("客户洞察及相关商机分析报告")
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.name = '黑体'

        # 报告周期
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"（{self.period}）")
        run.font.size = Pt(12)
        run.font.name = '宋体'

        # 报告日期
        current_date = datetime.now().strftime("%Y年%m月")
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(current_date)
        run.font.size = Pt(10.5)
        run.font.name = '宋体'

        # 分页
        self.doc.add_page_break()

    def _add_heading(self, text: str, level: int):
        """
        添加标题

        Args:
            text: 标题文本
            level: 标题层级（1-4）
        """
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.bold = True
        run.font.name = '黑体'

        # 根据层级设置字体大小
        font_sizes = {1: Pt(16), 2: Pt(14.5), 3: Pt(13), 4: Pt(12)}
        run.font.size = font_sizes.get(level, Pt(12))

        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(6)

        # 设置为标题样式
        heading_map = {
            1: self.doc.styles['Heading 1'],
            2: self.doc.styles['Heading 2'],
            3: self.doc.styles['Heading 3'],
            4: self.doc.styles['Heading 4']
        }
        para.style = heading_map.get(level, self.doc.styles['Heading 3'])

    def _add_paragraph(self, text: str):
        """
        添加段落

        Args:
            text: 段落文本
        """
        # 处理多行文本
        lines = text.split('\n')
        for line in lines:
            if line.strip():
                para = self.doc.add_paragraph()
                run = para.add_run(line)
                run.font.name = '宋体'
                run.font.size = Pt(10.5)
                para.paragraph_format.first_line_indent = Inches(0.35)
                para.paragraph_format.space_after = Pt(6)

    def _add_table(self, headers: List[str], rows: List[List[str]]):
        """
        添加表格

        Args:
            headers: 表头列表
            rows: 数据行列表
        """
        if not headers or not rows:
            return

        # 创建表格
        table = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.style = 'Table Grid'

        # 设置表头
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            cell = header_cells[i]
            cell.text = header
            # 设置表头样式
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.name = '黑体'
                    run.font.size = Pt(10)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 设置表头背景色（浅灰色）
            self._set_cell_background(cell, 'D9D9D9')

        # 填充数据
        for row_idx, row_data in enumerate(rows, start=1):
            row_cells = table.rows[row_idx].cells
            for col_idx, cell_data in enumerate(row_data):
                cell = row_cells[col_idx]
                cell.text = str(cell_data) if cell_data else '-'

                # 设置数据行样式
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = '宋体'
                        run.font.size = Pt(9)

                    # 根据数据类型设置对齐方式
                    if col_idx == 0:  # 第一列左对齐
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    elif '%' in str(cell_data):  # 百分比居中
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif self._is_numeric(cell_data):  # 数字右对齐
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

                # 设置单元格垂直居中
                from docx.enum.table import WD_ALIGN_VERTICAL
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                # 交替背景色
                if row_idx % 2 == 0:
                    self._set_cell_background(cell, 'F2F2F2')

        # 设置表格宽度
        table.autofit = False
        table.allow_autofit = False

    def _set_cell_background(self, cell, color_hex):
        """
        设置单元格背景色

        Args:
            cell: 单元格对象
            color_hex: 颜色十六进制值（如'D9D9D9'）
        """
        from docx.oxml import OxmlElement
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), color_hex)
        tcPr.append(shd)

    def _is_numeric(self, value: str) -> bool:
        """判断字符串是否为数字"""
        if not value:
            return False
        try:
            float(value.replace(',', '').replace('%', ''))
            return True
        except ValueError:
            return False

    def _process_chapter(self, chapter_data: Dict[str, Any]):
        """
        处理章节

        Args:
            chapter_data: 章节数据字典
        """
        chapter_title = chapter_data.get('chapter', '')
        subsections = chapter_data.get('subsections', [])

        # 处理标题中的【银行名称】占位符
        if '【银行名称】' in chapter_title and self.bank_name:
            chapter_title = chapter_title.replace('【银行名称】', self.bank_name)

        # 添加章节标题
        if chapter_title:
            self._add_heading(chapter_title, 1)

        # 递归处理子章节
        for subsection in subsections:
            self._process_subsection(subsection, 2)

    def _process_content(self, content_item: Dict[str, Any]):
        """
        处理内容项

        Args:
            content_item: 内容项字典，包含type字段
        """
        content_type = content_item.get('type')

        if content_type == 'heading':
            level = content_item.get('level', 1)
            text = content_item.get('text', '')
            self._add_heading(text, level)

        elif content_type == 'paragraph':
            text = content_item.get('text', '')
            self._add_paragraph(text)

        elif content_type == 'table':
            headers = content_item.get('headers', [])
            rows = content_item.get('rows', [])
            self._add_table(headers, rows)

    def _process_subsection(self, subsection: Dict[str, Any], default_level: int = 2):
        """
        递归处理子章节及其嵌套的子章节

        Args:
            subsection: 子章节数据字典
            default_level: 默认标题层级
        """
        # 获取子章节标题和层级
        subsection_title = subsection.get('title', '')
        heading_level = subsection.get('heading_level', default_level)

        # 处理标题中的【银行名称】占位符
        if '【银行名称】' in subsection_title and self.bank_name:
            subsection_title = subsection_title.replace('【银行名称】', self.bank_name)

        # 添加子章节标题
        if subsection_title:
            self._add_heading(subsection_title, heading_level)

        # 处理内容 - 支持content为对象或数组的情况
        content = subsection.get('content', [])
        if isinstance(content, dict):
            # 如果content是对象，转换为数组
            content = [content]

        # 处理内容列表
        for content_item in content:
            if isinstance(content_item, dict):
                self._process_content(content_item)

        # 处理直接定义在subsection级别的表格（Chapter 3和4的情况）
        if 'table' in subsection:
            table_data = subsection['table']
            headers = table_data.get('headers', [])
            rows = table_data.get('rows', [])
            if headers and rows:
                # 添加表格标题（如果有）
                table_title = table_data.get('table_title', '')
                table_number = table_data.get('table_number', '')
                if table_title:
                    title_text = f"{table_number} {table_title}" if table_number else table_title
                    self._add_heading(title_text, heading_level + 1)
                # 添加表格
                self._add_table(headers, rows)
                # 添加数据来源说明（如果有）
                data_source = table_data.get('data_source', '')
                if data_source:
                    self._add_paragraph(f"数据来源：{data_source}")

        # 递归处理嵌套的subsections
        nested_subsections = subsection.get('subsections', [])
        for nested_subsection in nested_subsections:
            self._process_subsection(nested_subsection, heading_level + 1)

    def generate(self, chapters: List[Dict[str, Any]], output_path: str):
        """
        生成Word文档

        Args:
            chapters: 章节数据列表
            output_path: 输出文件路径
        """
        # 设置样式
        self._setup_styles()

        # 添加封面页
        self._add_cover_page()

        # 处理所有章节
        for chapter_data in chapters:
            self._process_chapter(chapter_data)

        # 保存文档
        output_file = Path(output_path)
        output_file.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(output_file))

        print(f"报告已生成: {output_path}")
        print(f"  银行: {self.bank_name}")
        print(f"  周期: {self.period}")
        print(f"  章节数: {len(chapters)}")


def parse_arguments():
    """解析命令行参数"""
    parser = argparse.ArgumentParser(
        description='银行客户洞察报告生成器',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
使用示例:
  # 生成完整报告
  python generate_report.py \\
    --output-file "./宁波银行客户洞察报告_2025.docx" \\
    --bank-name "宁波银行" \\
    --period "2025年度" \\
    --json-content "filled_content.json"

  # 生成单章节报告
  python generate_report.py \\
    --output-file "./宁波银行基础信息_2025.docx" \\
    --bank-name "宁波银行" \\
    --period "2025年度" \\
    --json-content "chapter2_content.json"

JSON格式说明:
  {
    "chapters": [
      {
        "chapter": "一、基础信息",
        "subsections": [
          {
            "title": "（一）核心基础画像",
            "heading_level": 2,
            "content": [
              {
                "type": "heading",
                "level": 3,
                "text": "1. 主体及资质信息"
              },
              {
                "type": "paragraph",
                "text": "段落内容..."
              },
              {
                "type": "table",
                "headers": ["列1", "列2", "列3"],
                "rows": [
                  ["数据1", "数据2", "数据3"],
                  ["数据4", "数据5", "数据6"]
                ]
              }
            ]
          }
        ]
      }
    ]
  }
        '''
    )

    parser.add_argument(
        '--output-file',
        required=True,
        help='输出Word文件路径'
    )

    parser.add_argument(
        '--bank-name',
        required=True,
        help='银行名称'
    )

    parser.add_argument(
        '--period',
        required=True,
        help='报告周期（如"2025年度"）'
    )

    parser.add_argument(
        '--json-content',
        required=True,
        help='包含章节内容的JSON文件路径'
    )

    return parser.parse_args()


def main():
    """主函数"""
    args = parse_arguments()

    # 读取JSON内容
    try:
        with open(args.json_content, 'r', encoding='utf-8') as f:
            json_data = json.load(f)

        # JSON格式可能是直接包含chapters，或者是包装在其他对象中
        if isinstance(json_data, list):
            chapters = json_data
        elif isinstance(json_data, dict) and 'chapters' in json_data:
            chapters = json_data['chapters']
        elif isinstance(json_data, dict) and 'chapter' in json_data:
            # 单个章节
            chapters = [json_data]
        else:
            print(f"错误: JSON格式不正确，无法识别章节数据")
            sys.exit(1)

    except FileNotFoundError:
        print(f"错误: 文件不存在 - {args.json_content}")
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"错误: JSON解析失败 - {e}")
        sys.exit(1)

    # 生成报告
    generator = BankReportGenerator(args.bank_name, args.period)
    generator.generate(chapters, args.output_file)


if __name__ == '__main__':
    main()
